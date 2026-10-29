import datetime
import gzip
import json
import re
import shutil
from functools import lru_cache
from importlib.resources import files
from pathlib import Path

import httpx
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from neovax.io import digest, stable_id, write_json

PUBLISHED_COLUMNS = [
    "patient",
    "pool",
    "gene",
    "protein_change",
    "epitope_length",
    "hla",
    "mutant_epitope",
    "historical_mutant_ic50_nm",
    "wildtype_epitope",
    "historical_wildtype_ic50_nm",
    "immunizing_sequence",
    "imp_id",
    "ept_id",
    "gene_tpm",
]


def source_manifest():
    return json.loads(files("neovax").joinpath("resources/sources.json").read_text())


def fetch_data(destination, local_source=None, progress=lambda message: None):
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    for source in source_manifest():
        path = destination / source["file"]
        if path.exists() and digest(path) == source["sha256"]:
            progress(f"  Found and verified {path.name}; reusing it.")
            continue
        temporary = path.with_suffix(f"{path.suffix}.download")
        if local_source:
            progress(f"  Copying {path.name} from {local_source}")
            shutil.copyfile(Path(local_source) / source["file"], temporary)
        else:
            reason = "Checksum changed" if path.exists() else "Not found"
            progress(f"  {reason}: {path.name}. Downloading…")
            with httpx.stream("GET", source["url"], follow_redirects=True, timeout=120) as response:
                response.raise_for_status()
                with temporary.open("wb") as handle:
                    for chunk in response.iter_bytes():
                        handle.write(chunk)
        if digest(temporary) != source["sha256"]:
            temporary.unlink()
            raise ValueError(f"Hash mismatch for {source['file']}; the pinned reference may have changed")
        temporary.replace(path)
        progress(f"  Verified {path.name}.")
    verify_sources(destination)


@lru_cache(maxsize=4)
def verify_sources(directory):
    records = []
    for source in source_manifest():
        path = Path(directory) / source["file"]
        if not path.exists():
            raise FileNotFoundError(f"Missing {path}. Run neovax fetch-data first.")
        if digest(path) != source["sha256"]:
            raise ValueError(f"Input hash mismatch: {path}")
        records.append(source)
    return records


def hla_value(value):
    if isinstance(value, datetime.time):
        return f"{value.hour:02d}:{value.minute:02d}"
    if isinstance(value, datetime.timedelta):
        minutes = round(value.total_seconds() / 60)
        return f"{minutes // 60:02d}:{minutes % 60:02d}"
    return str(value).strip()


def normalize_hla(value):
    value = str(value).strip().replace("HLA-", "").replace("*", "")
    match = re.fullmatch(r"([AB])(\d{2}):(\d{2})", value)
    if not match:
        raise ValueError(f"Unsupported class I allele: {value}")
    return f"HLA-{match[1]}*{match[2]}:{match[3]}"


def result_state(value):
    if value is None or str(value).strip() == "":
        return "not reported"
    if str(value).strip().lower() in {"n.d.", "nd", "n.d"}:
        return "not determined"
    if value == 1 or str(value).strip().lower() in {"positive", "+"}:
        return "positive"
    if value == 0 or str(value).strip().lower() in {"negative", "-"}:
        return "negative"
    return "ambiguous"


def read_rows(path, sheet=None):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook.active
    title = worksheet.title
    rows = list(worksheet.values)
    workbook.close()
    return title, rows


@lru_cache(maxsize=4)
def load_reference(directory):
    with gzip.open(Path(directory) / "hg19_knownGeneTxPep.txt.gz", "rt") as handle:
        return dict(line.rstrip().split("\t", 1) for line in handle)


@lru_cache(maxsize=4)
def load_catalog(directory):
    directory = Path(directory)
    verify_sources(directory)
    cache_key = stable_id("catalog-v1", source_manifest())
    normalized = directory.parent / "outputs" / "cache" / f"normalized-{cache_key}.json"
    if normalized.exists():
        saved = json.loads(normalized.read_text())
        return (
            saved["variants"],
            {int(k): v for k, v in saved["alleles"].items()},
            saved["published"],
            saved["observations"],
        )
    variants, alleles, published, observations = [], {}, [], {}
    sheet, rows = read_rows(directory / "ott2017_table2.xlsx")
    for index, row in enumerate(rows[3:], 4):
        if row[0] not in range(1, 7):
            continue
        (
            patient,
            gene,
            transcript,
            consequence,
            _kind,
            build,
            chrom,
            start,
            end,
            strand,
            ref,
            a1,
            a2,
            _cdna,
            _codon,
            change,
        ) = row[:16]
        alternates = sorted({str(a) for a in (a1, a2) if a != ref})
        mutation_id = f"P{patient}:{build}:{chrom}:{start}:{end}:{ref}:{','.join(alternates)}"
        variants.append(
            {
                "variant_id": stable_id(mutation_id, transcript, change),
                "mutation_id": mutation_id,
                "patient": patient,
                "gene": gene,
                "transcript": transcript,
                "consequence": consequence,
                "protein_change": change,
                "build": build,
                "chromosome": chrom,
                "genomic_start": start,
                "genomic_end": end,
                "strand": strand,
                "reference_allele": ref,
                "tumor_alleles": [a1, a2],
                "source": f"ott2017_table2.xlsx#{sheet}!A{index}:P{index}",
            }
        )
    _, rows = read_rows(directory / "ott2017_table4.xlsx")
    for row in rows:
        if row[0] in range(1, 7):
            alleles[row[0]] = list(
                dict.fromkeys(f"HLA-{locus}*{hla_value(value)}" for locus, value in zip("AABB", row[1:5]))
            )

    sources = {s["file"]: s for s in source_manifest()}

    def observe(record, raw, file, sheet, cell, cell_type, assay, target, timepoint, stimulation):
        # Same biological observation occurs in several HLA rows and both publications.
        sequence = record.get("assay_sequence") or record.get("mutant_epitope") or record.get("imp_id")
        key = (record["patient"], sequence, cell_type, assay, target, timepoint, stimulation)
        source = {
            "file": file,
            "sha256": sources[file]["sha256"],
            "sheet": sheet,
            "cell": cell,
            "url": sources[file]["url"],
            "raw_value": raw,
        }
        state = result_state(raw)
        if key in observations:
            existing = observations[key]
            existing["sources"].append(source)
            if existing["result"] != state:
                existing["result"] = "ambiguous"
            return
        observations[key] = {
            "observation_id": stable_id(*key),
            "patient": record["patient"],
            "gene": record["gene"],
            "protein_change": record["protein_change"],
            "imp_id": record["imp_id"],
            "ept_id": record.get("ept_id"),
            "asp_id": record.get("asp_id"),
            "sequence": sequence,
            "cell_type": cell_type,
            "assay_type": assay,
            "target_cell": target,
            "timepoint": timepoint,
            "pre_stimulation": stimulation,
            "treatment_context": "post-vaccine; subsequent therapy may contribute"
            if timepoint != "week 16"
            else "vaccine week 16",
            "result": state,
            "value": None,
            "unit": None,
            "sources": [source],
            "restriction": "predicted HLA is not a confirmed experimental restriction",
        }

    sheet, rows = read_rows(directory / "ott2017_table5.xlsx")
    for index, row in enumerate(rows[4:], 5):
        if row[0] not in range(1, 7):
            continue
        record = dict(zip(PUBLISHED_COLUMNS, row[:14]))
        record["source"] = f"ott2017_table5.xlsx#{sheet}!A{index}:Q{index}"
        if record["hla"]:
            record["hla"] = normalize_hla(record["hla"])
        published.append(record)
        for col, assay, target in [
            (14, "peptide", "peptide-pulsed autologous APC"),
            (15, "minigene", "minigene-expressing autologous B cell"),
            (16, "tumor", "autologous tumor"),
        ]:
            observe(
                record,
                row[col],
                "ott2017_table5.xlsx",
                sheet,
                f"{get_column_letter(col + 1)}{index}",
                "CD8",
                assay,
                target,
                "week 16",
                "expanded",
            )
    sheet, rows = read_rows(directory / "hu2021_datasets.xlsx", "Suppl Dataset 4a. CD8+ T cells")
    for index, row in enumerate(rows[3:], 4):
        if row[0] not in range(1, 7):
            continue
        record = dict(zip(PUBLISHED_COLUMNS, row[:14]))
        for col, timepoint in [(14, "week 16"), (15, "year 3–4.5")]:
            observe(
                record,
                row[col],
                "hu2021_datasets.xlsx",
                sheet,
                f"{get_column_letter(col + 1)}{index}",
                "CD8",
                "peptide",
                "peptide-pulsed autologous APC",
                timepoint,
                "expanded",
            )
    sheet, rows = read_rows(directory / "hu2021_datasets.xlsx", "Suppl Dataset 4b. CD4+ T cells")
    for index, row in enumerate(rows[4:], 5):
        if row[0] not in range(1, 7):
            continue
        record = dict(
            zip(
                [
                    "patient",
                    "gene",
                    "protein_change",
                    "mutant_epitope",
                    "wildtype_epitope",
                    "imp_id",
                    "immunizing_sequence",
                    "asp_id",
                    "assay_sequence",
                ],
                row[:9],
            )
        )
        for col, assay, target, stimulation in [
            (9, "peptide", "peptide-pulsed autologous APC", "ex vivo"),
            (10, "peptide", "peptide-pulsed autologous APC", "expanded"),
            (11, "minigene", "minigene-expressing autologous B cell", "expanded"),
            (12, "HLA-DR blocking", "minigene-expressing autologous B cell", "expanded"),
            (13, "tumor", "autologous tumor", "expanded"),
            (14, "peptide", "peptide-pulsed autologous APC", "ex vivo"),
            (15, "peptide", "peptide-pulsed autologous APC", "expanded"),
        ]:
            observe(
                record,
                row[col],
                "hu2021_datasets.xlsx",
                sheet,
                f"{get_column_letter(col + 1)}{index}",
                "CD4",
                assay,
                target,
                "week 16" if col < 14 else "year 3–4.5",
                stimulation,
            )
    write_json(
        normalized,
        {
            "variants": variants,
            "alleles": alleles,
            "published": published,
            "observations": list(observations.values()),
        },
    )
    return variants, alleles, published, list(observations.values())


def patient_data(directory, patient):
    variants, alleles, published, assays = load_catalog(Path(directory))
    return (
        [dict(v) for v in variants if v["patient"] == patient],
        alleles[patient],
        [dict(v) for v in published if v["patient"] == patient],
        [dict(v) for v in assays if v["patient"] == patient],
    )


def actual_vaccine(published, variants, reconstructed):
    reconstructed_ids = {r["mutation_id"] for r in reconstructed}
    vaccine = {}
    for row in published:
        if row["imp_id"] in vaccine:
            continue
        matches = {
            v["mutation_id"]
            for v in variants
            if v["gene"] == row["gene"] and v["protein_change"] == row["protein_change"]
        }
        mutation_id = next(iter(matches)) if len(matches) == 1 else None
        vaccine[row["imp_id"]] = {
            k: row[k]
            for k in [
                "patient",
                "pool",
                "gene",
                "protein_change",
                "immunizing_sequence",
                "imp_id",
                "gene_tpm",
                "source",
            ]
        }
        vaccine[row["imp_id"]].update(
            mutation_id=mutation_id,
            target_join="audited gene + exact protein change → unique genomic mutation"
            if mutation_id
            else "ambiguous"
            if matches
            else "unresolved; no exact consequence match",
            reconstructable=mutation_id in reconstructed_ids,
        )
    return list(vaccine.values())
